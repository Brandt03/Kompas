"""Builds 'Oversigt alle år.xlsx': every category side by side per year, styled like Budget 2026.xlsx.
Usage: python3 yearly_overview.py <output dir>   (reads monthly.csv / balances.csv next to this file)"""
import csv, sys, datetime as dt
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.properties import PageSetupProperties

HERE = Path(__file__).parent
OUT = Path(sys.argv[1])
CUR = dt.date.today().year

data = defaultdict(float); group_of = {}
for y, m, grp, cat, kind, amt in csv.reader(open(HERE / "monthly.csv"), delimiter=";"):
    data[(int(y), cat)] += float(amt); group_of[cat] = grp
balance = {(int(y), int(m)): float(b) for y, m, b in csv.reader(open(HERE / "balances.csv"), delimiter=";")}
FIRST_YEAR, FIRST_MONTH = min(balance)   # the first month with a balance, i.e. when the account history starts
YEARS = list(range(FIRST_YEAR, CUR)) + [CUR]   # current year = year to date, kept out of totals and averages
FULL = YEARS[:-1]
months = {y: len({m for (yy, m) in balance if yy == y}) for y in YEARS}
year_end = {y: balance[max((yy, m) for (yy, m) in balance if yy == y)] for y in YEARS}
last_month = max(m for (y, m) in balance if y == CUR)

SECTIONS = [("INDTÆGTER", "Indtægter", -1, "I alt indtægter"), ("FASTE UDGIFTER", "Faste udgifter", 1, "I alt faste udgifter"),
            ("VARIABLE UDGIFTER", "Variable udgifter", 1, "I alt variable udgifter"), ("INDFLYTNING", "Indflytning", 1, "I alt indflytning"),
            ("OPSPARING OG INVESTERING", "Opsparing og investering", 1, "I alt opsparing")]
NAVY, BLUE, SECTION, TOTAL, INPUT, GREY, YTD = "1F3864", "8EAADB", "D9E2F3", "E2EFDA", "FFF2CC", "595959", "F2F2F2"
NF = '#,##0;[Red]\\-#,##0;\\–'
thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
def font(sz=10, b=False, color="000000", i=False): return Font(name="Arial", size=sz, bold=b, color=color, italic=i)
def fill(c): return PatternFill("solid", fgColor=c)
def cell(ws, ref, value, *, b=False, fl=None, nf=None, align=None, color="000000", border=True, sz=10):
    c = ws[ref]; c.value = value; c.font = font(sz, b, color)
    if fl: c.fill = fill(fl)
    if nf: c.number_format = nf
    if align: c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=False)
    if border: c.border = BOX
    return c
def title(ws, text, sub, to):
    ws.merge_cells(f"B2:{to}2"); ws["B2"] = text; ws["B2"].font = font(16, True, "FFFFFF"); ws["B2"].fill = fill(NAVY)
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1); ws.row_dimensions[2].height = 30
    ws["B3"] = sub; ws["B3"].font = font(10, color=GREY); ws.sheet_view.showGridLines = False; ws.column_dimensions["A"].width = 2

wb = Workbook(); ws = wb.active; ws.title = "Alle år"
YC = {y: L(3 + i) for i, y in enumerate(YEARS)}           # C..N
FIRST, LAST = YC[FULL[0]], YC[FULL[-1]]                   # full years only
SUMC, AVGC = L(3 + len(YEARS)), L(4 + len(YEARS))         # O, P
title(ws, f"ALLE ÅR – {FIRST_YEAR} TIL {CUR-1} SIDE OM SIDE",
      f"Faktiske tal fra Sure pr. kalenderår. {CUR} er år-til-dato (jan.–{['jan','feb','mar','apr','maj','jun','jul','aug','sep','okt','nov','dec'][last_month-1]}.) og indgår ikke i 'I alt' og gennemsnit.", AVGC)

r = 5
heads = [("B", "Post")] + [(YC[y], f"{y}" if y != CUR else f"{CUR} (ÅTD)") for y in YEARS] + [(SUMC, f"I alt {FIRST_YEAR}–{str(CUR-1)[2:]}"), (AVGC, "Gns./år")]
for col, h in heads: cell(ws, f"{col}{r}", h, b=True, fl=BLUE, color="FFFFFF", align="left" if col == "B" else "center")
r += 1; mrow = r
cell(ws, f"B{r}", "Måneder med data", align="left", color=GREY)
for y in YEARS: cell(ws, f"{YC[y]}{r}", months[y], nf="0", align="right", color="0000FF")
cell(ws, f"{SUMC}{r}", f"=SUM({FIRST}{r}:{LAST}{r})", nf="0", align="right"); cell(ws, f"{AVGC}{r}", None)

totals, rows_of = {}, {}
for sec, grp, sign, tl in SECTIONS:
    names = sorted([c for c, g in group_of.items() if g == grp and any(round(data.get((y, c), 0)) for y in YEARS)], key=str.casefold)
    r += 2
    cell(ws, f"B{r}", sec, b=True, fl=SECTION, color=NAVY, sz=11, align="left")
    for col in [YC[y] for y in YEARS] + [SUMC, AVGC]: cell(ws, f"{col}{r}", None, fl=SECTION)
    start = r + 1
    for n in names:
        r += 1; cell(ws, f"B{r}", n, align="left")
        for y in YEARS:
            v = round(sign * data.get((y, n), 0), 2)
            cell(ws, f"{YC[y]}{r}", v, fl=YTD if y == CUR else INPUT, nf=NF, align="right")
        cell(ws, f"{SUMC}{r}", f"=SUM({FIRST}{r}:{LAST}{r})", b=True, nf=NF, align="right")
        cell(ws, f"{AVGC}{r}", f"={SUMC}{r}/{len(FULL)}", nf=NF, align="right")
    rows_of[grp] = (start, r)
    r += 1; totals[grp] = r
    cell(ws, f"B{r}", tl, b=True, fl=TOTAL, align="left")
    for col in [YC[y] for y in YEARS] + [SUMC, AVGC]:
        cell(ws, f"{col}{r}", f"=SUM({col}{start}:{col}{r-1})", b=True, fl=TOTAL, nf=NF, align="right")

T = lambda col, g: f"{col}{totals[g]}"
r += 2; res = r
cell(ws, f"B{r}", "ÅRETS RESULTAT", b=True, fl=TOTAL, sz=11, align="left")
for col in [YC[y] for y in YEARS] + [SUMC, AVGC]:
    cell(ws, f"{col}{r}", "=" + T(col, "Indtægter") + "".join(f"-{T(col, g)}" for g in ["Faste udgifter", "Variable udgifter", "Indflytning", "Opsparing og investering"]),
         b=True, fl=TOTAL, nf=NF, align="right")
extra = [
    ("Forbrug pr. måned (faste + variable)", lambda c: f"=IFERROR(({T(c,'Faste udgifter')}+{T(c,'Variable udgifter')})/{c}{mrow},\"\")", NF),
    ("Indtægt pr. måned", lambda c: f"=IFERROR({T(c,'Indtægter')}/{c}{mrow},\"\")", NF),
    ("Faste udgifter i % af indtægt", lambda c: f"=IFERROR({T(c,'Faste udgifter')}/{T(c,'Indtægter')},\"\")", "0%;[Red]-0%;\\–"),
    ("Opsparingsgrad", lambda c: f"=IFERROR(({T(c,'Opsparing og investering')}+MAX(0,{c}{res}))/{T(c,'Indtægter')},\"\")", "0%;[Red]-0%;\\–"),
]
for label, f, nf in extra:  # noqa
    r += 1; cell(ws, f"B{r}", label, align="left")
    for col in [YC[y] for y in YEARS] + [SUMC]: cell(ws, f"{col}{r}", f(col), nf=nf, align="right")
    cell(ws, f"{AVGC}{r}", None)
r += 1; bal = r
cell(ws, f"B{r}", "Saldo på kontiene ved årets udgang", b=True, align="left")
for y in YEARS: cell(ws, f"{YC[y]}{r}", round(year_end[y], 2), fl=YTD if y == CUR else INPUT, nf=NF, align="right")
cell(ws, f"{SUMC}{r}", None); cell(ws, f"{AVGC}{r}", None)
r += 2
MONTHS_LONG = ["januar", "februar", "marts", "april", "maj", "juni", "juli", "august", "september", "oktober", "november", "december"]
for line in ["Alle beløb er i kroner. Gule felter er hentet fra Sure og kan rettes – totaler, nøgletal og grafer regner selv videre."] + (
             [f"{FIRST_YEAR} dækker kun {MONTHS_LONG[FIRST_MONTH-1]}–december, fordi kontoen starter der. Sammenlign derfor hellere 'pr. måned'-rækkerne end årstotalerne for {FIRST_YEAR}."]
             if FIRST_MONTH > 1 else []) + [
             "Overførsler mellem egne konti er holdt ude. Refusioner trækker fra i den kategori, de hører til."]:
    ws[f"B{r}"] = line; ws[f"B{r}"].font = font(9, color=GREY); r += 1
ws.column_dimensions["B"].width = 36
for y in YEARS: ws.column_dimensions[YC[y]].width = 10.5
ws.column_dimensions[SUMC].width = 13; ws.column_dimensions[AVGC].width = 10
ws.freeze_panes = "C6"
ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0

# ---- Grafer ---------------------------------------------------------------------
g = wb.create_sheet("Grafer")
title(g, "GRAFER – ÅR FOR ÅR", f"Kun hele år ({FIRST_YEAR}–{CUR-1}). Datagrundlaget står i kolonne S og frem.", "Q")
A = "'Alle år'"
top_var = sorted(range(*[rows_of["Variable udgifter"][0], rows_of["Variable udgifter"][1] + 1]),
                 key=lambda rr: -sum(data.get((y, ws[f"B{rr}"].value), 0) for y in FULL))[:6]
hdrs = ["År", "Indtægter", "Udgifter", "Resultat", "Saldo", "Forbrug/md."] + [ws[f"B{rr}"].value for rr in top_var]
for i, h in enumerate(hdrs):
    c = g.cell(row=4, column=19 + i, value=h); c.font = font(9, True, "FFFFFF"); c.fill = fill(BLUE)
for k, y in enumerate(FULL):
    rr, c = 5 + k, YC[y]
    vals = [str(y), f"={A}!{T(c,'Indtægter')}", f"={A}!{T(c,'Faste udgifter')}+{A}!{T(c,'Variable udgifter')}+{A}!{T(c,'Indflytning')}",
            f"={A}!{c}{res}", f"={A}!{c}{bal}", f"={A}!{c}{res+1}"] + [f"={A}!{c}{t}" for t in top_var]
    for i, v in enumerate(vals):
        cell_ = g.cell(row=rr, column=19 + i, value=v); cell_.font = font(9); cell_.number_format = NF
for i in range(len(hdrs)): g.column_dimensions[L(19 + i)].width = 13
for col in "BCDEFGHIJKLMNOPQ": g.column_dimensions[col].width = 9
n = len(FULL); cats = Reference(g, min_col=19, min_row=5, max_row=4 + n)

def chart(ch, t, anchor, w=13.5, h=8):
    ch.title = t; ch.width = w; ch.height = h; ch.style = 10; ch.y_axis.numFmt = "#,##0"; g.add_chart(ch, anchor); return ch
def paint(s, col, line=False):
    if line: s.graphicalProperties.line.solidFill = col; s.graphicalProperties.line.width = 28000; s.smooth = False
    else: s.graphicalProperties.solidFill = col; s.graphicalProperties.line.solidFill = col

b = BarChart(); b.add_data(Reference(g, min_col=20, max_col=21, min_row=4, max_row=4 + n), titles_from_data=True); b.set_categories(cats)
paint(b.series[0], "70AD47"); paint(b.series[1], "ED7D31"); b.gapWidth = 60; b.y_axis.majorGridlines = None
chart(b, "Indtægter og udgifter pr. år", "B5")
rb = BarChart(); rb.add_data(Reference(g, min_col=22, min_row=4, max_row=4 + n), titles_from_data=True); rb.set_categories(cats)
paint(rb.series[0], "1F3864"); rb.series[0].invertIfNegative = False; rb.legend = None; rb.gapWidth = 60; rb.x_axis.tickLblPos = "low"
chart(rb, "Årets resultat", "J5")
ln = LineChart(); ln.add_data(Reference(g, min_col=23, min_row=4, max_row=4 + n), titles_from_data=True); ln.set_categories(cats)
paint(ln.series[0], "8EAADB", line=True); ln.legend = None
chart(ln, "Saldo på kontiene ved årets udgang", "B22")
fm = LineChart(); fm.add_data(Reference(g, min_col=24, min_row=4, max_row=4 + n), titles_from_data=True); fm.set_categories(cats)
paint(fm.series[0], "ED7D31", line=True); fm.legend = None
chart(fm, "Forbrug pr. måned (faste + variable)", "J22")
st = BarChart(); st.grouping = "stacked"; st.overlap = 100
st.add_data(Reference(g, min_col=25, max_col=24 + len(top_var), min_row=4, max_row=4 + n), titles_from_data=True); st.set_categories(cats)
for s, col in zip(st.series, ["1F3864", "8EAADB", "ED7D31", "70AD47", "FFC000", "A5A5A5"]): paint(s, col)
st.gapWidth = 50; st.y_axis.majorGridlines = None
chart(st, "De 6 største variable udgifter – år for år", "B39", 27.5, 9.5)
g["B59"] = f"{CUR} er ikke med i graferne, fordi året ikke er slut. Se kolonnen '{CUR} (ÅTD)' i fanen Alle år."
g["B59"].font = font(9, color=GREY)

OUT.mkdir(parents=True, exist_ok=True)
wb.calculation.fullCalcOnLoad = True
p = OUT / "Oversigt alle år.xlsx"; tmp = p.with_name(f".~{p.name}"); wb.save(tmp); tmp.replace(p); print(p)
