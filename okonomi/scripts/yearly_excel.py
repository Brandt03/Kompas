"""Builds one 'Regnskab <år>.xlsx' per year from Sure's monthly category totals, styled like Budget 2026.xlsx.
Usage: yearly_excel.py <output dir> <year> [<year> ...] [--budget "Budget <år>.xlsx"]
With --budget, the monthly budget (Månedsbudget, column C) is mapped onto Sure's categories and shown next to
the latest complete month. The budget file is only read, never written."""
import argparse, csv, datetime as dt, warnings
warnings.filterwarnings("ignore", module="openpyxl")  # Budget file has Excel-only extensions; we only read it
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, AreaChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.worksheet.properties import PageSetupProperties

HERE = Path(__file__).parent
ap = argparse.ArgumentParser()
ap.add_argument("out"); ap.add_argument("years", nargs="+", type=int); ap.add_argument("--budget")
ARGS = ap.parse_args()
OUT = Path(ARGS.out)

# Budget 2026 'Månedsbudget' labels (prefix) -> Sure category
BUDGET_MAP = [
    ("SU", "SU"), ("Løn", "Løn"), ("Boligstøtte", "Boligstøtte"), ("Tilskud fra familie", "Tilskud fra familie"),
    ("Andre indtægter", "Andre indtægter"),
    *[(k, "Husleje") for k in ("Husleje", "Kontingent", "Varme", "Vand", "El", "Internet")],  # paid as one sum to the landlord
    ("Fællesudgifter", "Fællesudgifter (fælleskøkken)"),
    *[(k, "Forsikringer") for k in ("Indboforsikring", "Sundhedsforsikring", "Ulykkesforsikring")],
    ("Mobilabonnement", "Mobilabonnement"), ("A-kasse", "A-kasse og fagforening"), ("Transport", "Transport"),
    ("Streaming", "Streaming og abonnementer"), ("Sport", "Sport og træning"),
    ("Afdrag på lån", "Afdrag på lån"), ("Studiebøger", "Studiebøger og materialer"),
    *[(k, k) for k in ("Dagligvarer", "Take-away og mad ude", "Café, bar og byture", "Tøj og sko", "Personlig pleje og frisør",
                        "Medicin, læge og tandlæge", "Husholdning og rengøring", "Gaver og fester", "Fritid, hobby og kultur",
                        "Ferie og rejser", "Uforudsete udgifter")],
    *[(k, "Opsparing og investering") for k in ("Nødopsparing", "Målopsparing", "Investering")],
]

def read_budget(path):
    ws = load_workbook(path, data_only=True)["Månedsbudget"]
    out = defaultdict(float)
    for r in range(1, ws.max_row + 1):
        label, val = ws[f"B{r}"].value, ws[f"C{r}"].value
        if not isinstance(label, str) or not isinstance(val, (int, float)) or label.startswith("I alt"):
            continue
        # longest matching prefix wins ("SU-lån" and "SU efter skat" both -> SU; "El" must not catch "Elektronik")
        hits = [(k, c) for k, c in BUDGET_MAP if label == k or label.startswith(k + " ") or label.startswith(k + "-") or label.startswith(k + " (")]
        if hits:
            out[max(hits, key=lambda h: len(h[0]))[1]] += val
    return dict(out)
BUDGET = read_budget(ARGS.budget) if ARGS.budget and Path(ARGS.budget).exists() else None

# ---- data -------------------------------------------------------------------
data = defaultdict(float)            # (year, month, category) -> kr (positive = expense/outflow in Sure terms)
group_of = {}
for y, m, grp, cat, kind, amt in csv.reader(open(HERE / "monthly.csv"), delimiter=";"):
    data[(int(y), int(m), cat)] += float(amt)
    group_of[cat] = grp
balance = {(int(y), int(m)): float(b) for y, m, b in csv.reader(open(HERE / "balances.csv"), delimiter=";")}
first_month = min((y, m) for (y, m) in balance)

def cats(group):
    return sorted([c for c, g in group_of.items() if g == group], key=str.casefold)

SECTIONS = [  # (title, sure group, sign, total label)
    ("INDTÆGTER", "Indtægter", -1, "I alt indtægter"),
    ("FASTE UDGIFTER", "Faste udgifter", 1, "I alt faste udgifter"),
    ("VARIABLE UDGIFTER", "Variable udgifter", 1, "I alt variable udgifter"),
    ("INDFLYTNING", "Indflytning", 1, "I alt indflytning"),
    ("OPSPARING OG INVESTERING", "Opsparing og investering", 1, "I alt opsparing"),
]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "Maj", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dec"]

# ---- styles (lifted from Budget 2026.xlsx) ----------------------------------
NAVY, BLUE, SECTION, TOTAL, INPUT, GREY = "1F3864", "8EAADB", "D9E2F3", "E2EFDA", "FFF2CC", "595959"
NF = '#,##0;[Red]\\-#,##0;\\–'
thin = Side(style="thin", color="BFBFBF")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
def font(sz=10, b=False, color="000000"): return Font(name="Arial", size=sz, bold=b, color=color)
def fill(c): return PatternFill("solid", fgColor=c)

def title(ws, text, sub, width_to="Q"):
    ws.merge_cells(f"B2:{width_to}2")
    ws["B2"] = text; ws["B2"].font = font(16, True, "FFFFFF"); ws["B2"].fill = fill(NAVY)
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 30
    ws["B3"] = sub; ws["B3"].font = font(10, color=GREY)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2

def cell(ws, ref, value, *, b=False, fl=None, nf=None, align=None, color="000000", border=True, sz=10):
    c = ws[ref]; c.value = value; c.font = font(sz, b, color)
    if fl: c.fill = fill(fl)
    if nf: c.number_format = nf
    if align: c.alignment = Alignment(horizontal=align, vertical="center")
    if border: c.border = BOX
    return c

# Årsoversigt columns: B post, C..N months, O total, P avg/md, Q share of group
MC = [chr(ord("C") + i) for i in range(12)]
TOT, AVG, SHARE = "O", "P", "Q"

def build(year, budget=None):
    active = [m for m in range(1, 13) if (year, m) >= first_month and (year, m) in balance]
    current = year == dt.date.today().year
    done = [m for m in active if not (current and m == dt.date.today().month)]   # complete months
    wb = Workbook(); ws = wb.active; ws.title = "Årsoversigt"
    wb.calculation.fullCalcOnLoad = True
    stamp = f" Opdateret automatisk {dt.datetime.now():%d-%m-%Y kl. %H:%M} – ret kategorier i Sure, ikke her." if current else ""
    title(ws, f"ÅRSOVERSIGT {year} – FAKTISK FORBRUG" + (" (ÅR TIL DATO)" if current else ""),
          "Tallene er hentet fra Sure. Overførsler mellem egne konti er holdt ude." + stamp,
          "T" if budget else "Q")
    ws["B4"] = "Måneder med data"; ws["B4"].font = font(9, color=GREY)
    cell(ws, "C4", len(active), nf="0", color="0000FF", border=False, sz=9, align="left")
    ws["D4"] = "← bruges til 'Gns./md.'" + (f" (kontoen starter {MONTHS[active[0]-1].lower()}. {year})" if active[0] > 1 else "")
    ws["D4"].font = font(9, color=GREY)

    r = 5
    BM = MC[done[-1] - 1] if budget and done else None   # column of the latest complete month
    extra_h = [("R", "Budget/md."), ("S", f"{MONTHS[done[-1]-1]} faktisk" if BM else "Faktisk"), ("T", "Afvigelse")] if budget else []
    for col, h in [("B", "Post")] + list(zip(MC, MONTHS)) + [(TOT, "I alt året"), (AVG, "Gns./md."), (SHARE, "Andel")] + extra_h:
        cell(ws, f"{col}{r}", h, b=True, fl=BLUE, color="FFFFFF", align="left" if col == "B" else "center")
    ws.row_dimensions[r].height = 19.5
    totals, rows_of = {}, {}
    for sec, grp, sign, tlabel in SECTIONS:
        names = [n for n in cats(grp) if any(round(data.get((year, m, n), 0)) for m in range(1, 13)) or (budget or {}).get(n)]
        if not names:
            totals[grp] = None
            continue
        r += 1
        cell(ws, f"B{r}", sec, b=True, fl=SECTION, color=NAVY, sz=11, align="left")
        for col in MC + [TOT, AVG, SHARE] + ([c for c, _ in extra_h]): cell(ws, f"{col}{r}", None, fl=SECTION)
        start = r + 1
        for name in names:
            r += 1
            cell(ws, f"B{r}", name, align="left")
            for i, col in enumerate(MC):
                v = data.get((year, i + 1, name))
                cell(ws, f"{col}{r}", round(sign * v, 2) if v else 0, fl=INPUT, nf=NF, align="right")
            cell(ws, f"{TOT}{r}", f"=SUM(C{r}:N{r})", b=True, nf=NF, align="right")
            cell(ws, f"{AVG}{r}", f"=IF($C$4=0,\"\",{TOT}{r}/$C$4)", nf=NF, align="right")
            if budget:
                bv = budget.get(name)
                cell(ws, f"R{r}", round(bv, 2) if bv is not None else None, color="008000", nf=NF, align="right")
                cell(ws, f"S{r}", f"={BM}{r}" if BM else None, nf=NF, align="right")
                # positive = better than budget (more income / less spending)
                cell(ws, f"T{r}", (f"=IF(R{r}=\"\",\"\",{'S'+str(r)+'-R'+str(r) if sign < 0 else 'R'+str(r)+'-S'+str(r)})") if BM else None, nf=NF, align="right")
        end = r
        rows_of[grp] = (start, end)
        r += 1
        cell(ws, f"B{r}", tlabel, b=True, fl=TOTAL, align="left")
        for col in MC + [TOT, AVG]:
            cell(ws, f"{col}{r}", f"=SUM({col}{start}:{col}{end})", b=True, fl=TOTAL, nf=NF, align="right")
        cell(ws, f"{SHARE}{r}", None, fl=TOTAL)
        for col in [c for c, _ in extra_h]:
            f = f"=SUM({col}{start}:{col}{end})" if col != "T" else (f"=S{r}-R{r}" if sign < 0 else f"=R{r}-S{r}")
            cell(ws, f"{col}{r}", f if (col != "T" or BM) else None, b=True, fl=TOTAL, nf=NF, align="right")
        for rr in range(start, end + 1):
            cell(ws, f"{SHARE}{rr}", f"=IF(${TOT}${r}=0,\"\",{TOT}{rr}/${TOT}${r})", nf="0%;[Red]-0%;\\–", align="right")
        totals[grp] = r
        r += 1
    r += 1
    res = r
    def T(col, grp):  # reference to a section total, or 0 when the section is empty this year
        return f"{col}{totals[grp]}" if totals.get(grp) else "0"
    def TA(col, grp):  # same, from the Grafer sheet
        return f"'{A}'!{col}{totals[grp]}" if totals.get(grp) else "0"
    A = "Årsoversigt"
    cell(ws, f"B{r}", "MÅNEDENS RESULTAT", b=True, fl=TOTAL, sz=11, align="left")
    for col in MC + [TOT, AVG]:
        cell(ws, f"{col}{r}", f"={T(col,'Indtægter')}-{T(col,'Faste udgifter')}-{T(col,'Variable udgifter')}-{T(col,'Indflytning')}-{T(col,'Opsparing og investering')}", b=True, fl=TOTAL, nf=NF, align="right")
    cell(ws, f"{SHARE}{r}", None, fl=TOTAL)
    if budget:
        for col in ("R", "S"):
            cell(ws, f"{col}{r}", f"={T(col,'Indtægter')}-{T(col,'Faste udgifter')}-{T(col,'Variable udgifter')}-{T(col,'Indflytning')}-{T(col,'Opsparing og investering')}", b=True, fl=TOTAL, nf=NF, align="right")
        cell(ws, f"T{r}", f"=S{r}-R{r}" if BM else None, b=True, fl=TOTAL, nf=NF, align="right")
    r += 1
    last_entry = max([m for m in range(1, 13) if any(k[0] == year and k[1] == m for k in data)] + active[-1:])
    cell(ws, f"B{r}", "Resultat akkumuleret over året", b=True, align="left")
    for i, col in enumerate(MC):
        f = (f"=C{res}" if i == 0 else f"={MC[i-1]}{r}+{col}{res}") if i < last_entry else None
        cell(ws, f"{col}{r}", f, nf=NF, align="right")
    last_entry = max([m for m in range(1, 13) if any(k[0] == year and k[1] == m for k in data)] + active[-1:])
    LASTC = MC[active[-1] - 1]   # latest month with a balance (saldo)
    cell(ws, f"{TOT}{r}", f"={MC[last_entry - 1]}{r}", nf=NF, align="right")   # incl. future-dated scheduled payments
    r += 1; bal = r
    cell(ws, f"B{r}", "Saldo på kontiene ved månedens udgang", b=True, align="left")
    for i, col in enumerate(MC):
        v = balance.get((year, i + 1))
        cell(ws, f"{col}{r}", round(v, 2) if v is not None else None, fl=INPUT if v is not None else None, nf=NF, align="right")
    cell(ws, f"{TOT}{r}", f"={LASTC}{r}", nf=NF, align="right")

    # Nøgletal
    r += 2
    cell(ws, f"B{r}", "NØGLETAL", b=True, fl=SECTION, color=NAVY, sz=11, align="left", border=False)
    kt = [
        ("Faste udgifter i % af indtægt", f"=IFERROR({T(TOT,'Faste udgifter')}/{T(TOT,'Indtægter')},\"\")", "0%", "Det du er bundet til hver måned."),
        ("Opsparingsgrad", f"=IFERROR(({T(TOT,'Opsparing og investering')}+MAX(0,{TOT}{res}))/{T(TOT,'Indtægter')},\"\")", "0%", "Opsparing + overskud i forhold til indtægten."),
        ("Forbrug pr. dag (faste + variable)", f"=IFERROR(({T(TOT,'Faste udgifter')}+{T(TOT,'Variable udgifter')})/($C$4*30.4),\"\")", NF, "Gennemsnit over månederne med data."),
        ("Største variable post", f"=INDEX(B{rows_of['Variable udgifter'][0]}:B{rows_of['Variable udgifter'][1]},MATCH(MAX({TOT}{rows_of['Variable udgifter'][0]}:{TOT}{rows_of['Variable udgifter'][1]}),{TOT}{rows_of['Variable udgifter'][0]}:{TOT}{rows_of['Variable udgifter'][1]},0))", None, None),
        ("Saldo ændret over året", f"=IFERROR({LASTC}{bal}-INDEX(C{bal}:N{bal},MATCH(TRUE,INDEX(C{bal}:N{bal}<>\"\",0),0)),\"\")", NF, "Fra første til seneste måned med data."),
    ]
    for label, f, nf, note in kt:
        r += 1
        cell(ws, f"B{r}", label, align="left")
        cell(ws, f"C{r}", f, nf=nf, align="right", b=True)
        ws.merge_cells(f"C{r}:D{r}")
        if note: ws[f"E{r}"] = note; ws[f"E{r}"].font = font(9, color=GREY)
    r += 2
    for line in [
        "Alle beløb er i kroner. Indtægter står som plus, udgifter som plus under deres gruppe; refusioner trækker fra i den kategori, de hører til (derfor kan en udgift blive negativ i en måned).",
        "Gule felter er hentet fra Sure og kan rettes. Totaler, gennemsnit, resultat og graferne regner selv videre.",
        "Kun kategorier med bevægelser i året er med. Tomme sektioner (fx Indflytning) er udeladt.",
    ] + ([
        f"Budget/md. (grøn) hentes fra fanen Månedsbudget i {Path(ARGS.budget).name}. Husleje, kontingent, varme, vand, el og internet er lagt sammen, fordi de betales samlet til udlejeren; forsikringerne ligeså.",
        "Afvigelse: plus = bedre end budgettet (mere ind eller mindre ud), minus = dårligere. Sammenlignes med seneste hele måned.",
    ] if budget else []):
        ws[f"B{r}"] = line; ws[f"B{r}"].font = font(9, color=GREY); r += 1

    ws.column_dimensions["B"].width = 34
    for col in MC: ws.column_dimensions[col].width = 9.5
    ws.column_dimensions[TOT].width = 12; ws.column_dimensions[AVG].width = 10; ws.column_dimensions[SHARE].width = 8
    if budget:
        for col in "RST": ws.column_dimensions[col].width = 11
    ws.freeze_panes = "C6"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0

    # ---- Grafer --------------------------------------------------------------
    g = wb.create_sheet("Grafer")
    title(g, f"GRAFER OG OVERBLIK {year}", "Graferne regner selv videre ud fra Årsoversigt. Datagrundlaget står i kolonne S–Z til højre.", "Q")
    A = "Årsoversigt"
    def hdr(ref, text): cell(g, ref, text, b=True, fl=BLUE, color="FFFFFF", border=False)
    # Month table S..W
    hdr("S4", "Måned"); hdr("T4", "Indtægter"); hdr("U4", "Udgifter"); hdr("V4", "Resultat"); hdr("W4", "Saldo")
    for i, col in enumerate(MC):
        rr = 5 + i
        g[f"S{rr}"] = MONTHS[i]
        g[f"T{rr}"] = f"={TA(col,'Indtægter')}"
        g[f"U{rr}"] = f"={TA(col,'Faste udgifter')}+{TA(col,'Variable udgifter')}+{TA(col,'Indflytning')}"
        g[f"V{rr}"] = f"='{A}'!{col}{res}"
        if (year, i + 1) in balance: g[f"W{rr}"] = f"='{A}'!{col}{bal}"
    # Split of the year
    hdr("S19", "Fordeling af året"); hdr("T19", "Beløb")
    split = [("Faste udgifter", f"={TA(TOT,'Faste udgifter')}"), ("Variable udgifter", f"={TA(TOT,'Variable udgifter')}"),
             ("Indflytning", f"=MAX(0,{TA(TOT,'Indflytning')})"), ("Opsparing", f"=MAX(0,{TA(TOT,'Opsparing og investering')})"), ("Til overs", f"=MAX(0,'{A}'!{TOT}{res})")]
    for i, (k, f) in enumerate(split):
        g[f"S{20+i}"] = k; g[f"T{20+i}"] = f
    # Per-post tables Y..Z (income, fixed, variable)
    def post_table(top, label, grp):
        if grp not in rows_of: return None
        hdr(f"Y{top}", label); hdr(f"Z{top}", "Året")
        s, e = rows_of[grp]
        for k, rr in enumerate(range(s, e + 1)):
            g[f"Y{top+1+k}"] = f"='{A}'!B{rr}"; g[f"Z{top+1+k}"] = f"='{A}'!{TOT}{rr}"
        return top + 1, top + 1 + (e - s)
    inc = post_table(4, "Indtægter", "Indtægter")
    fix = post_table(inc[1] + 3, "Faste udgifter", "Faste udgifter")
    var = post_table((fix or inc)[1] + 3, "Variable udgifter", "Variable udgifter")
    for rr in range(4, (var or fix or inc)[1] + 1):
        for col in "STUVWYZ":
            c = g[f"{col}{rr}"]
            if c.value is not None and not c.font.b: c.font = font(9); c.number_format = NF
    for col in "BCDEFGHIJKLMNOPQ": g.column_dimensions[col].width = 9
    for col, w in {"R": 3, "S": 18, "T": 11, "U": 11, "V": 11, "W": 11, "X": 3, "Y": 30, "Z": 11}.items():
        g.column_dimensions[col].width = w

    def style(ch, t, w=13.5, h=8):
        ch.title = t; ch.width = w; ch.height = h; ch.style = 10
        return ch
    def color(series, hexcol, line=False):
        if line:
            series.graphicalProperties.line.solidFill = hexcol; series.graphicalProperties.line.width = 28000; series.smooth = False
        else:
            series.graphicalProperties.solidFill = hexcol; series.graphicalProperties.line.solidFill = hexcol

    d = style(DoughnutChart(), "Hvor blev pengene af?")
    d.add_data(Reference(g, min_col=20, min_row=19, max_row=24), titles_from_data=True)
    d.set_categories(Reference(g, min_col=19, min_row=20, max_row=24))
    for i, c in enumerate(["1F3864", "ED7D31", "C00000", "70AD47", "8EAADB"]):
        pt = DataPoint(idx=i); pt.graphicalProperties.solidFill = c; d.series[0].dPt.append(pt)
    d.dataLabels = DataLabelList(); d.dataLabels.showPercent = True
    for k in ("showVal", "showSerName", "showCatName", "showLegendKey", "showLeaderLines"): setattr(d.dataLabels, k, False)
    d.holeSize = 55
    g.add_chart(d, "B5")

    b = style(BarChart(), "Indtægter og udgifter måned for måned")
    b.add_data(Reference(g, min_col=20, max_col=21, min_row=4, max_row=16), titles_from_data=True)
    b.set_categories(Reference(g, min_col=19, min_row=5, max_row=16))
    color(b.series[0], "70AD47"); color(b.series[1], "ED7D31"); b.y_axis.numFmt = "#,##0"; b.gapWidth = 60
    b.y_axis.majorGridlines = None
    g.add_chart(b, "J5")

    def hbar(t, rng, anchor, col):
        if not rng: return
        ch = style(BarChart(), t, 13.5, max(8, 0.55 * (rng[1] - rng[0] + 1) + 2)); ch.type = "bar"
        ch.add_data(Reference(g, min_col=26, min_row=rng[0], max_row=rng[1]), titles_from_data=False)
        ch.set_categories(Reference(g, min_col=25, min_row=rng[0], max_row=rng[1]))
        ch.legend = None; color(ch.series[0], col); ch.y_axis.numFmt = "#,##0"; ch.x_axis.scaling.orientation = "maxMin"
        ch.y_axis.majorGridlines = None; ch.gapWidth = 40
        g.add_chart(ch, anchor)
    hbar("Indtægter – hvor kom pengene fra?", inc, "B22", "70AD47")
    hbar("Faste udgifter – post for post", fix, "J22", "1F3864")
    hbar("Variable udgifter – post for post", var, "B40", "ED7D31")

    ln = style(LineChart(), "Resultat måned for måned")
    ln.add_data(Reference(g, min_col=22, min_row=4, max_row=16), titles_from_data=True)
    ln.set_categories(Reference(g, min_col=19, min_row=5, max_row=16))
    color(ln.series[0], "1F3864", line=True); ln.y_axis.numFmt = "#,##0"; ln.legend = None
    g.add_chart(ln, "J40")

    ar = style(AreaChart(), "Saldo på kontiene ved månedens udgang", 27.5)
    ar.add_data(Reference(g, min_col=23, min_row=4, max_row=16), titles_from_data=True)
    ar.set_categories(Reference(g, min_col=19, min_row=5, max_row=16))
    color(ar.series[0], "8EAADB"); ar.y_axis.numFmt = "#,##0"; ar.legend = None
    g.add_chart(ar, "B61")
    g["B79"] = "Månederne før kontoen startede står tomme i saldo-grafen. Negative måneder i 'Resultat' betyder, at der gik flere penge ud end ind."
    g["B79"].font = font(9, color=GREY)

    wb.properties.creator = "Sure → Excel"
    path = OUT / f"Regnskab {year}.xlsx"
    tmp = path.with_name(f".~{path.name}")
    wb.save(tmp); tmp.replace(path)   # atomic swap, so OneDrive/Excel never see a half-written file
    return path

OUT.mkdir(parents=True, exist_ok=True)
for y in ARGS.years:
    print(build(y, BUDGET if y == dt.date.today().year else None))
